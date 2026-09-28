"""Resume upload and retrieval endpoints."""
import secrets
from datetime import UTC, datetime
from io import BytesIO

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pypdf import PdfReader

from app.api.dependencies import current_user
from app.schemas.resume import ResumeResponse, ResumeUploadResponse
from app.services.resume_parser import extract_resume_profile
from app.services.store import RESUMES

router = APIRouter(prefix="/resume", tags=["resume"])
MAX_UPLOAD_BYTES = 8 * 1024 * 1024


def _extract_text(filename: str, data: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix in {"txt", "md"}:
            return data.decode("utf-8-sig")
        if suffix == "pdf":
            return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
        if suffix == "docx":
            document = Document(BytesIO(data))
            blocks = [p.text for p in document.paragraphs if p.text.strip()]
            blocks.extend(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                          for table in document.tables for row in table.rows)
            blocks.extend(p.text for section in document.sections
                          for p in (*section.header.paragraphs, *section.footer.paragraphs) if p.text.strip())
            return "\n".join(blocks)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Unable to parse the uploaded resume") from exc
    raise HTTPException(status_code=415, detail="Supported resume formats are PDF, DOCX, TXT, and MD")


@router.post("/upload", status_code=201, response_model=ResumeUploadResponse)
async def upload_resume(file: UploadFile = File(...), user: dict = Depends(current_user)) -> dict:
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Resume must be 8 MB or smaller")
    text = _extract_text(file.filename or "resume.txt", data).strip()
    if not text:
        raise HTTPException(status_code=422, detail="No readable text was found in the resume")
    profile = extract_resume_profile(text)
    resume_id = secrets.token_urlsafe(16)
    record = {"id": resume_id, "user_id": user["id"], "filename": file.filename,
              "parsed_text": text, "profile": profile, "skills": profile["skills"],
              "created_at": datetime.now(UTC).isoformat()}
    RESUMES[resume_id] = record
    return {k: v for k, v in record.items() if k != "parsed_text"} | {"text_characters": len(text)}


@router.get("/{resume_id}", response_model=ResumeResponse)
def get_resume(resume_id: str, user: dict = Depends(current_user)) -> dict:
    record = RESUMES.get(resume_id)
    if record is None or record["user_id"] != user["id"]:
        raise HTTPException(status_code=404, detail="Resume not found")
    return record
